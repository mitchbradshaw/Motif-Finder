"""
kinds.py
========
One ``KindSpec`` per registrable kind (``docs/DATA_REGISTRATION.md``):
conventional roots, how to scan them, the checks, the table a registration
becomes, and where the UI shows it. Adding a kind = one ``KindSpec`` here,
one test, one row in the doc.

    recording             DATA/derived/channels/<stem>/CH<n>.npy + manifest.json  -> recordings (one row per channel)
    raw                   DATA/raw/*.mat|*.csv                                     -> derived to a recording, then recordings
    model                 MODELS/*.pth, DATA/derived/models/*.joblib               -> registered_artifacts(kind='model')
    window_matrix         Results/Preprocessing/window_matrix/wm_v1_*.npz, MATRICES/*.csv
    matrix_profile        Results/Detection/matrix_profile/mp_v2_*.npz
    window_set            DATA/derived/window_sets/<name>/{windows.npz, manifest.json}
    encoding              DATA/derived/encodings/*.npz                             -> encodings
    drop_motif_store      <dir>/{events.csv, snippets.npz, manifest.json}          (Working/Detection/drop_motifs/store.py)
    catalogue_spreadsheet DATA/catalogue/*.xlsx
    hpc_result            HPC/results/<job>/{<job>.json, artifacts, manifest.json?}

Binding is by CONTENT: an artifact names its recording by (source_file,
channel); the local ``recordings`` id is looked up, never trusted from the
file (a cluster's ids are not ours). Checks never raise — a broken file is a
failed check with the reason in ``detail``.
"""
from __future__ import annotations

import csv
import glob
import json
import os
import re
import shutil
import sqlite3
from dataclasses import dataclass
from typing import Callable

import numpy as np

from .core import ACTOR, Candidate, Report, now_iso, plain_insert, portable, read_sidecar, sidecar_path
from .excerpt import R_THRESHOLD, find_excerpt, is_excerpt
from Working.units import UNITS, parse_units

HELD_OUT_FILE = "M4_aug_concat_fs1.mat"
HELD_OUT_STEM = HELD_OUT_FILE[:-4]
EXCERPT_MAX_SAMPLES = 200_000      # an excerpt is short; a same-length pair is a copy, not an excerpt
FS_RATIO_TOL = 1e-6

_CH_FILE = re.compile(r"^(?:.*_)?CH(\d+)\.npy$")
_FS_SUFFIX = re.compile(r"_fs(\d+(?:\.\d+)?)$", re.IGNORECASE)
_MP_NAME = re.compile(r"^mp_v2_(?P<stem>.+)_CH(?P<ch>\d+)_WIN(?P<win>[\d.]+)min(?:_span(?P<a>\d+)-(?P<b>\d+))?\.npz$")
_WM_NAME = re.compile(r"^wm_v1_(?P<stem>.+)_CH(?P<ch>\d+)_WIN(?P<win>[\d.]+)min_STEP(?P<step>[\d.]+)pct\.npz$")
_HEX8 = re.compile(r"(?<![0-9a-f])([0-9a-f]{8})(?![0-9a-f])")


@dataclass
class KindSpec:
    name: str
    label: str
    roots: list
    table: str
    ui: str
    naming: str
    scan: Callable
    check: Callable
    register: Callable
    registered_ids: Callable
    unregister: Callable
    list_registered: Callable
    enrich: Callable | None = None      # fill a registered candidate's facts from its rows

    def describe(self) -> dict:
        return {"name": self.name, "label": self.label, "roots": list(self.roots), "table": self.table, "ui": self.ui, "naming": self.naming}


# ================================================================== helpers ==

def _abs(p: str) -> str:
    return os.path.abspath(p).replace("\\", "/").lower()


def _same_path(a: str, b: str) -> bool:
    return _abs(a) == _abs(b)


def _rows(conn, sql, args=()):
    cur = conn.execute(sql, args)
    cols = [d[0] for d in cur.description]
    return [dict(zip(cols, r)) for r in cur.fetchall()]


def _recording_rows(conn, source_file=None, channel=None):
    sql = "SELECT * FROM recordings WHERE active = 1"
    args = []
    if source_file is not None:
        sql += " AND source_file = ?"; args.append(source_file)
    if channel is not None:
        sql += " AND channel = ?"; args.append(int(channel))
    return _rows(conn, sql + " ORDER BY source_file, channel", tuple(args))


def _bind_recording(conn, rep: Report, source_file, channel, stem=None) -> dict | None:
    """The local recordings row for (source_file, channel) — by content."""
    if conn is None:
        rep.add("recording", False, "no database connection to bind the recording")
        return None
    rows = _recording_rows(conn, source_file, channel) if source_file else []
    if not rows and stem:
        rows = _recording_rows(conn, f"{stem}.mat", channel)
    if not rows:
        rep.add("recording", False, f"no registered recording for source_file={source_file!r} channel={channel} (stem {stem!r}); register the recording first")
        return None
    row = rows[0]
    rep.add("recording", True, f"bound to recordings id {row['id']} ({row['source_file']} CH{row['channel']}, fs {row['fs']}, {row['n_samples']:,} samples)")
    rep.facts["recording_id"] = row["id"]
    rep.facts["source_file"] = row["source_file"]
    rep.facts["channel"] = row["channel"]
    return row


def _scalar(z, key, default=None):
    if key not in z.files:
        return default
    v = z[key]
    try:
        v = v.item() if getattr(v, "shape", None) == () else v
    except Exception:
        pass
    if isinstance(v, (bytes, np.bytes_)):
        v = v.decode("utf-8", "replace")
    if isinstance(v, np.generic):
        v = v.item()
    return v


def _npz_member_shape(path: str, member: str):
    """Shape of one array in an .npz from its .npy header alone (no data read)."""
    import zipfile
    with zipfile.ZipFile(path) as zf:
        name = member + ".npy"
        if name not in zf.namelist():
            return None
        with zf.open(name) as fh:
            version = np.lib.format.read_magic(fh)
            shape, _fortran, _dtype = np.lib.format._read_array_header(fh, version)
            return tuple(shape)


def _npz_keys(z, rep: Report, required) -> bool:
    missing = [k for k in required if k not in z.files]
    return rep.add("keys", not missing, f"missing keys {missing}" if missing else f"keys {list(z.files)[:12]}{'…' if len(z.files) > 12 else ''}")


def _finite(arr, rep: Report, label="values") -> bool:
    a = np.asarray(arr)
    if a.dtype.kind not in "fc":
        return rep.add("finite", True, f"{label}: integer dtype {a.dtype}")
    n_nan = int(np.isnan(a).sum())
    n_inf = int(np.isinf(a).sum())
    if n_inf:
        rep.warn(f"{label}: {n_inf:,} infinite values")
    return rep.add("finite", n_nan == 0, f"{label}: {n_nan:,} NaN" if n_nan else f"{label}: all {a.size:,} values finite" + (f" ({n_inf} inf)" if n_inf else ""))


def _fs_matches(rep: Report, fs_file, row) -> bool:
    if fs_file is None:
        return rep.add("fs", False, "the file carries no fs")
    ok = abs(float(fs_file) - float(row["fs"])) <= 1e-9 * max(1.0, float(row["fs"]))
    return rep.add("fs", ok, f"fs {fs_file} vs recording {row['fs']}" + ("" if ok else " — disagree"))


def _span_inside(rep: Report, span, row) -> bool:
    a, b = span
    n = int(row["n_samples"])
    ok = 0 <= a < b <= n
    rep.facts["span"] = [int(a), int(b)]
    return rep.add("span", ok, f"span [{a:,}, {b:,}) inside the recording's {n:,} samples" if ok else f"span [{a:,}, {b:,}) is not inside [0, {n:,})")


def _json_dumps(v):
    return json.dumps(v, default=str) if v is not None else None


# -------------------------------------------------- registered_artifacts kinds --

def _art_registered_ids(conn, cand: Candidate, **kw) -> list:
    rows = _rows(conn, "SELECT id, path FROM registered_artifacts WHERE kind = ? AND active = 1", (cand.kind,))
    return [r["id"] for r in rows if _same_path(r["path"], cand.path)]


def _art_register(conn, cand: Candidate, rep: Report, provenance, actor, writer, **kw):
    f = rep.facts
    prov = dict(provenance or {})
    span = f.get("span")
    row = {
        "kind": cand.kind, "path": cand.path, "name": cand.name, "manifest_path": None,
        "recording_id": f.get("recording_id"), "channel": f.get("channel"),
        "span_start": span[0] if span else None, "span_end": span[1] if span else None,
        "fs": f.get("fs"), "params_json": _json_dumps(f.get("parameters", {})),
        "producer": prov.get("producer") or f.get("producer"), "sha1": rep.sha1,
        "checks_json": _json_dumps([{"name": c.name, "ok": c.ok, "detail": c.detail} for c in rep.checks]),
        "warnings_json": _json_dumps(list(rep.warnings)), "created_at": now_iso(), "actor": actor, "active": 1,
    }
    rid = writer(conn, "registered_artifacts", row)
    return rid, row


def _art_unregister(conn, row_id, actor):
    conn.execute("UPDATE registered_artifacts SET active = 0 WHERE id = ?", (int(row_id),))
    return {"table": "registered_artifacts", "id": int(row_id), "active": 0}


def _art_list(kind):
    def _list(conn, sidecar_root=None, **kw):
        out = []
        for r in _rows(conn, "SELECT * FROM registered_artifacts WHERE kind = ? AND active = 1 ORDER BY id", (kind,)):
            r["params"] = json.loads(r["params_json"]) if r.get("params_json") else {}
            r["warnings"] = json.loads(r["warnings_json"]) if r.get("warnings_json") else []
            r["checks"] = json.loads(r["checks_json"]) if r.get("checks_json") else []
            r["exists"] = os.path.exists(r["path"])
            r["manifest"] = read_sidecar(r["path"], sidecar_root)
            r["bytes"] = _tree_bytes(r["path"]) if r["exists"] else None
            out.append(r)
        return out
    return _list


def _tree_bytes(path: str) -> int:
    if os.path.isfile(path):
        return os.path.getsize(path)
    total = 0
    for root, _d, files in os.walk(path):
        for fn in files:
            try:
                total += os.path.getsize(os.path.join(root, fn))
            except OSError:
                pass
    return total


def _files(roots, exts, recursive=False):
    out = []
    for root in roots:
        if not os.path.isdir(root):
            continue
        if recursive:
            for r, dirs, files in os.walk(root):
                dirs[:] = sorted(d for d in dirs if not d.startswith((".", "_")))
                out += [os.path.join(r, f) for f in sorted(files) if f.lower().endswith(exts)]
        else:
            out += [os.path.join(root, f) for f in sorted(os.listdir(root)) if f.lower().endswith(exts) and os.path.isfile(os.path.join(root, f))]
    return [portable(p) for p in out]


# ================================================================ recording ==

def _channel_files(d: str):
    found = []
    for fn in sorted(os.listdir(d)):
        m = _CH_FILE.match(fn)
        if m:
            found.append((int(m.group(1)), fn))
    found.sort()
    return found


def _read_manifest(d: str):
    p = os.path.join(d, "manifest.json")
    if not os.path.isfile(p):
        return None
    try:
        with open(p, encoding="utf-8") as f:
            return json.load(f)
    except Exception as e:
        return {"error": f"{type(e).__name__}: {e}"}


def _scan_recording(roots, conn, **kw):
    cands = []
    for root in roots:
        if not os.path.isdir(root):
            continue
        for name in sorted(os.listdir(root)):
            d = os.path.join(root, name)
            if name.startswith(".") or not os.path.isdir(d):
                continue
            chans = _channel_files(d)
            if not chans:
                continue
            man = _read_manifest(d) or {}
            warnings = []
            if "error" in man:
                warnings.append(f"manifest.json unreadable: {man['error']}")
            fs = man.get("fs")
            n = man.get("n_samples_per_channel")
            if n is None:
                try:
                    n = int(np.load(os.path.join(d, chans[0][1]), mmap_mode="r").shape[0])
                except Exception as e:
                    warnings.append(f"{chans[0][1]} unreadable: {type(e).__name__}: {e}")
            fs_note = man.get("fs_note")
            fs_source = "read"
            if fs is None:
                warnings.append("fs unknown: no manifest.json carries it — supply the sampling rate at registration (it is recorded as inferred)")
            if fs_note and "not uniform" in fs_note:
                warnings.append(f"sampling is not uniform: {fs_note}")
            if fs_note and "inferred" in fs_note:
                warnings.append(f"fs is inferred: {fs_note}")
                fs_source = "inferred"
            units_text = man.get("units")
            units = parse_units(units_text)
            warnings.extend(_units_warnings(units_text, units))
            facts = {"stem": name, "dir": portable(d), "source_file": man.get("source_file") or f"{name}.mat", "source_file_from_manifest": bool(man.get("source_file")),
                     "fs": fs, "fs_source": fs_source if fs is not None else None, "fs_note": fs_note, "n_channels": len(chans), "n_samples": n,
                     "dtype": man.get("dtype"), "time_base": man.get("time_base"), "units": units, "units_text": units_text, "imported_at": man.get("imported_at"),
                     "raw_file": man.get("raw_file"), "has_manifest": bool(man) and "error" not in man, "channel_files": [fn for _i, fn in chans],
                     "duration_h": (n / fs / 3600.0) if (fs and n) else None, "held_out": man.get("source_file") == HELD_OUT_FILE or name == HELD_OUT_STEM}
            if not man.get("source_file"):
                warnings.append(f"no manifest.json: the recording is labelled '{name}.mat' after its directory; pass source_file to override")
            cands.append(Candidate(kind="recording", path=portable(d), name=name, facts=facts, warnings=warnings))
    return cands


UNITS_UNDECLARED = ("units undeclared: manifest.json declares no unit for these samples — supply units "
                    f"({' or '.join(UNITS[:2])}) at registration, or declare it later in Settings › Datasets; until then "
                    "every page prints this recording's numbers as unit undeclared, never as mV")


def _units_warnings(units_text, units) -> list:
    """The scan's words about the unit: nothing when it is declared, a warning
    when it is absent, a different one when it is there but unreadable."""
    if units is not None:
        return []
    if units_text is None or not str(units_text).strip():
        return [UNITS_UNDECLARED]
    return [f"units unreadable: manifest.json says {units_text!r}, which is not one of {', '.join(UNITS)} — "
            "supply units at registration"]


def _recording_ids_for_dir(conn, d: str, source_file: str | None = None) -> list:
    rows = _rows(conn, "SELECT id, npy_path, source_file FROM recordings WHERE active = 1")
    ids = [r["id"] for r in rows if _same_path(os.path.dirname(r["npy_path"]), d)]
    if source_file:
        ids += [r["id"] for r in rows if r["source_file"] == source_file and r["id"] not in ids]
    return sorted(ids)


def _recording_registered_ids(conn, cand: Candidate, **kw) -> list:
    return _recording_ids_for_dir(conn, cand.path, cand.facts.get("source_file"))


def _enrich_recording(conn, cand: Candidate, **kw) -> None:
    """A registered directory's facts come from its rows: the pre-standard
    directories carry no manifest.json, and the row is the truth anyway."""
    if not cand.registered_ids:
        return
    r = conn.execute("SELECT source_file, fs, n_samples, fs_source, units, units_note FROM recordings WHERE id = ?", (cand.registered_ids[0],)).fetchone()
    if r is None:
        return
    sf, fs, n, fs_source = r[0], r[1], r[2], r[3]
    cand.facts.update({"source_file": sf, "fs": fs, "n_samples": n, "fs_source": fs_source or "unrecorded",
                       "duration_h": n / fs / 3600.0 if fs else None, "facts_from": "recordings rows",
                       "units": r[4], "units_note": r[5]})
    # the row is the truth: a unit declared on it answers the manifest's silence
    if r[4]:
        cand.warnings = [w for w in cand.warnings if not w.startswith(("units undeclared", "units unreadable"))]


def _load_channels(d, files):
    return [np.load(os.path.join(d, fn), mmap_mode="r") for fn in files]


def _check_recording(cand: Candidate, conn, rep: Report, ov: dict, **kw):
    d = cand.path
    f = rep.facts
    files = f.get("channel_files") or [fn for _i, fn in _channel_files(d)]
    if ov.get("source_file"):
        f["source_file"] = str(ov["source_file"])
    # readable + shape
    lengths, bad = [], []
    arrays = []
    for fn in files:
        try:
            a = np.load(os.path.join(d, fn), mmap_mode="r")
            if a.ndim != 1:
                bad.append(f"{fn}: {a.ndim}-D, expected 1-D"); continue
            if a.dtype.kind not in "fiu":
                bad.append(f"{fn}: dtype {a.dtype}"); continue
            head, tail = np.asarray(a[:1000], dtype=np.float64), np.asarray(a[-1000:], dtype=np.float64)
            if not (np.isfinite(head).all() and np.isfinite(tail).all()):
                bad.append(f"{fn}: non-finite samples at the head or tail")
            lengths.append(int(a.shape[0])); arrays.append(a)
        except Exception as e:
            bad.append(f"{fn}: {type(e).__name__}: {e}")
    rep.add("readable", not bad, "; ".join(bad) if bad else f"{len(files)} channel files load (1-D, mmap)")
    if lengths:
        same = len(set(lengths)) == 1
        n_manifest = f.get("n_samples")
        ok = same and (n_manifest is None or n_manifest == lengths[0])
        rep.add("shape", ok, f"{len(lengths)} channels × {lengths[0]:,} samples" if ok else
                (f"channel lengths differ: {sorted(set(lengths))}" if not same else f"manifest says {n_manifest:,} samples, files hold {lengths[0]:,}"))
        if same:
            f["n_samples"] = lengths[0]
    f["n_channels"] = len(files)
    # fs
    fs = ov.get("fs", f.get("fs"))
    if fs is None:
        rep.add("fs", False, "sampling rate unknown — supply fs (recorded as inferred)")
    else:
        try:
            fs = float(fs)
            ok = fs > 0
        except (TypeError, ValueError):
            ok = False
        rep.add("fs", ok, f"fs {fs} Hz ({'supplied' if 'fs' in ov else 'from manifest'})" if ok else f"bad fs {fs!r}")
        if ok:
            f["fs"] = fs
            if "fs" in ov:
                f["fs_source"] = ov.get("fs_source", "inferred")
                # the scan's "fs unknown" warning is answered by the supplied value: it must not follow the row
                rep.warnings = [w for w in rep.warnings if not w.startswith("fs unknown")]
                if f["fs_source"] == "inferred":
                    rep.warn(f"fs {fs} Hz was supplied at registration, not read from the file: recorded as inferred")
                else:
                    rep.warn(f"fs {fs} Hz supplied at registration as read from the raw file's time vector (see the sidecar's notes)")
            elif f.get("fs_source") is None:
                f["fs_source"] = "read"
            if f.get("n_samples"):
                f["duration_h"] = f["n_samples"] / fs / 3600.0
    # units: flagged in words when undeclared (the file is still readable data), refused when unreadable
    if ov.get("units") is not None and str(ov.get("units")).strip():
        u = parse_units(ov["units"])
        rep.add("units", u is not None, f"units {u} (supplied)" if u else
                f"units {ov['units']!r} not recognised — one of {', '.join(UNITS)}")
        if u:
            f["units"] = u
            f["units_note"] = "declared at registration"
            rep.warnings = [w for w in rep.warnings if not w.startswith(("units undeclared", "units unreadable"))]
    elif f.get("units"):
        rep.add("units", True, f"units {f['units']} (from manifest.json)")
        f.setdefault("units_note", f"from manifest.json: {f.get('units_text') or f['units']}")
    # held out
    held = f.get("source_file") == HELD_OUT_FILE or cand.name == HELD_OUT_STEM
    rep.add("held_out", not held, f"{HELD_OUT_FILE} is held out (spec §0 D6): it is never registered through the interface" if held else "not the held-out recording")
    # excerpt links (both directions)
    if conn is not None and arrays and f.get("fs") and rep.ok:
        _excerpt_checks(conn, rep, arrays, float(f["fs"]), int(f["n_samples"]), ov, cand)
    else:
        rep.add("excerpt", True, "excerpt check skipped (earlier check failed or no database)")


def _excerpt_checks(conn, rep: Report, arrays, fs, n, ov, cand: Candidate):
    rows = _rows(conn, "SELECT id, source_file, channel, fs, n_samples, npy_path FROM recordings WHERE active = 1 AND source_file != ?", (rep.facts.get("source_file"),))
    found_of = None
    for r in rows:
        if not os.path.isfile(r["npy_path"]):
            continue
        fs_r, n_r = float(r["fs"]), int(r["n_samples"])
        if fs >= fs_r:                       # the registered row may be an excerpt of the candidate
            ratio = fs / fs_r
            d = int(round(ratio))
            if d < 1 or abs(ratio - d) > FS_RATIO_TOL or n_r > EXCERPT_MAX_SAMPLES or n_r * d > n:
                continue
            short = np.asarray(np.load(r["npy_path"], mmap_mode="r"), dtype=np.float64)
            for ci, a in enumerate(arrays):
                hit = find_excerpt(short, np.asarray(a, dtype=np.float64), d)
                if is_excerpt(hit):
                    ex = {"recording_id": r["id"], "source_file": r["source_file"], "channel": int(r["channel"]), "candidate_channel": ci,
                          "offset": int(hit["offset"]), "decimation": d, "r": round(hit["r"], 5), "neighbours": [round(x, 4) for x in hit["neighbours"]]}
                    rep.excerpts.append(ex)
                    rep.warn(f"registered {r['source_file']} CH{r['channel']} (id {r['id']}) is a {d}:1 excerpt of this recording's CH{ci} at sample {hit['offset']:,} (r = {hit['r']:.4f}); it will be linked, its id kept")
                    break
        else:                                # the candidate may be an excerpt of the registered row
            ratio = fs_r / fs
            d = int(round(ratio))
            if d < 1 or abs(ratio - d) > FS_RATIO_TOL or n > EXCERPT_MAX_SAMPLES or n * d > n_r or found_of:
                continue
            long = np.asarray(np.load(r["npy_path"], mmap_mode="r"), dtype=np.float64)
            for ci, a in enumerate(arrays):
                hit = find_excerpt(np.asarray(a, dtype=np.float64), long, d)
                if is_excerpt(hit):
                    found_of = {"recording_id": r["id"], "source_file": r["source_file"], "channel": int(r["channel"]), "candidate_channel": ci,
                                "offset": int(hit["offset"]), "decimation": d, "r": round(hit["r"], 5), "name": os.path.basename(os.path.dirname(r["npy_path"]))}
                    break
    if found_of:
        rep.excerpt_of = found_of
        stem = found_of["name"]
        msg = (f"this recording is a {found_of['decimation']}:1 excerpt of {stem} CH{found_of['channel']} (id {found_of['recording_id']}) "
               f"at sample {found_of['offset']:,} (r = {found_of['r']:.4f})")
        if ov.get("allow_excerpt"):
            rep.warn(msg + " — registered as an excerpt, linked to its parent")
            rep.add("excerpt", True, msg + "; allowed by the researcher")
        else:
            rep.add("excerpt", False, msg + " — a subset of a registered recording is an excerpt, not a new recording; pass allow_excerpt to register it linked to its parent")
    else:
        rep.add("excerpt", True, f"not an excerpt of any registered recording (r > {R_THRESHOLD} with a one-sample peak against every comparable channel)")


def _register_recording(conn, cand: Candidate, rep: Report, provenance, actor, writer, **kw):
    f = rep.facts
    d = cand.path
    files = f.get("channel_files") or [fn for _i, fn in _channel_files(d)]
    ids = []
    row0 = None
    ex = rep.excerpt_of
    for fn in files:
        ch = int(_CH_FILE.match(fn).group(1))
        row = {"source_file": f["source_file"], "channel": ch, "fs": float(f["fs"]), "n_samples": int(f["n_samples"]), "global_offset": 0,
               "npy_path": portable(os.path.join(d, fn)), "notes": (provenance or {}).get("notes"),
               "fs_source": f.get("fs_source") or "read", "registered_at": now_iso(), "registered_by": actor,
               "warnings_json": _json_dumps(list(rep.warnings)), "active": 1,
               "units": f.get("units"), "units_note": f.get("units_note") if f.get("units") else None}
        if ex and ex.get("candidate_channel") == ch:
            row.update({"parent_recording_id": ex["recording_id"], "parent_offset": ex["offset"], "decimation": ex["decimation"]})
        rid = writer(conn, "recordings", row)
        ids.append(rid)
        row0 = row0 or row
    for e in rep.excerpts:
        parent_id = ids[e["candidate_channel"]] if e["candidate_channel"] < len(ids) else ids[0]
        conn.execute("UPDATE recordings SET parent_recording_id = ?, parent_offset = ?, decimation = ? WHERE id = ?",
                     (parent_id, int(e["offset"]), int(e["decimation"]), int(e["recording_id"])))
    f["registered_ids"] = ids
    return ids[0], row0


def _unregister_recording(conn, row_id, actor):
    row = conn.execute("SELECT source_file FROM recordings WHERE id = ?", (int(row_id),)).fetchone()
    if row is None:
        raise KeyError(f"no recordings row {row_id}")
    sf = row[0]
    conn.execute("UPDATE recordings SET active = 0 WHERE source_file = ?", (sf,))
    n = conn.execute("SELECT COUNT(*) FROM recordings WHERE source_file = ? AND active = 0", (sf,)).fetchone()[0]
    return {"table": "recordings", "id": int(row_id), "source_file": sf, "channels": n, "active": 0}


def declare_units(conn, source_file: str, units: str, note: str | None = None, actor: str = ACTOR) -> dict:
    """Declare the unit an already-registered recording's samples are stored in.

    Every channel row of ``source_file`` takes the unit together — a recording
    is one export in one unit. The note says who declared it and on what
    evidence; the schema backfill never overwrites a row that carries one.
    Raises ``ValueError`` for a unit it cannot read, ``KeyError`` for a file
    with no rows. Does not commit (the caller audits and commits)."""
    u = parse_units(units)
    if u is None:
        raise ValueError(f"units {units!r} not recognised — one of {', '.join(UNITS)}")
    n = conn.execute("SELECT COUNT(*) FROM recordings WHERE source_file = ?", (source_file,)).fetchone()[0]
    if not n:
        raise KeyError(f"no recordings rows for {source_file!r}")
    text = (note or "").strip() or f"declared by {actor} on {now_iso()[:10]}"
    conn.execute("UPDATE recordings SET units = ?, units_note = ? WHERE source_file = ?", (u, text, source_file))
    return {"source_file": source_file, "units": u, "units_note": text, "channels": int(n)}


def _list_recordings(conn, sidecar_root=None, **kw):
    rows = _rows(conn, "SELECT * FROM recordings WHERE active = 1 ORDER BY source_file, channel")
    by = {}
    for r in rows:
        by.setdefault(r["source_file"], []).append(r)
    out = []
    for sf, chans in by.items():
        c0 = chans[0]
        d = os.path.dirname(c0["npy_path"])
        stem = os.path.basename(d) or sf.rsplit(".", 1)[0]
        parent = None
        if c0.get("parent_recording_id"):
            p = conn.execute("SELECT id, source_file, channel, npy_path FROM recordings WHERE id = ?", (c0["parent_recording_id"],)).fetchone()
            if p is not None:
                parent = {"recording_id": p[0], "source_file": p[1], "channel": p[2], "name": os.path.basename(os.path.dirname(p[3])),
                          "offset": c0.get("parent_offset"), "decimation": c0.get("decimation")}
        out.append({
            "kind": "recording", "id": c0["id"], "ids": [c["id"] for c in chans], "name": stem, "source_file": sf, "dir": portable(d),
            "n_channels": len(chans), "fs": c0["fs"], "fs_source": c0.get("fs_source") or "unrecorded", "n_samples": c0["n_samples"],
            "units": c0.get("units"), "units_note": c0.get("units_note"),
            "duration_h": c0["n_samples"] / c0["fs"] / 3600.0 if c0["fs"] else None, "held_out": sf == HELD_OUT_FILE,
            "warnings": json.loads(c0["warnings_json"]) if c0.get("warnings_json") else [], "excerpt_of": parent,
            "registered_at": c0.get("registered_at"), "registered_by": c0.get("registered_by"),
            "channels": [{"id": c["id"], "channel": c["channel"], "npy_path": c["npy_path"], "exists": os.path.isfile(c["npy_path"]),
                          "parent_recording_id": c.get("parent_recording_id")} for c in chans],
            "npy_exists": all(os.path.isfile(c["npy_path"]) for c in chans),
            "manifest": read_sidecar(d, sidecar_root) if d else None,
        })
    return out


# ====================================================================== raw ==

def _mat5_variables(path: str):
    """`scipy.io.whosmat` without its one weakness: a MATLAB opaque object (a
    `duration`, a table) has no dims and makes it raise TypeError for the whole
    file (F2B.mat, L_LM_Jul_26_J_raw.mat, M101_t.mat). Same header walk, one
    try/except per variable; an opaque variable is listed with shape None."""
    from scipy.io.matlab._mio5 import MatFile5Reader
    from scipy.io.matlab._mio5_params import mclass_info
    out = []
    with open(path, "rb") as fh:
        rd = MatFile5Reader(fh)
        rd.mat_stream.seek(0)
        rd.initialize_read()
        rd.read_file_header()
        while not rd.end_of_stream():
            hdr, nxt = rd.read_var_header()
            name = "None" if hdr.name is None else hdr.name.decode("latin1")
            try:
                shape = list(rd._matrix_reader.shape_from_header(hdr))
            except Exception:
                shape = None
            info = "logical" if getattr(hdr, "is_logical", False) else mclass_info.get(hdr.mclass, "opaque")
            out.append((name, shape, info))
            rd.mat_stream.seek(nxt)
    return out


def _mat_variables(path: str):
    """[(name, shape, dtype)] without loading the data (v5 via the header walk, v7.3 via h5py)."""
    import scipy.io
    try:
        try:
            return [(n, list(s), str(t)) for n, s, t in scipy.io.whosmat(path)], "v5"
        except TypeError:
            return _mat5_variables(path), "v5"
    except NotImplementedError:
        import h5py
        out = []
        with h5py.File(path, "r") as f:
            for k in f.keys():
                if k.startswith("#"):
                    continue
                ds = f[k]
                out.append((k, list(getattr(ds, "shape", [])), str(getattr(ds, "dtype", "?"))))
        return out, "v7.3"


def _layout_guess(variables):
    """Which variable holds the data and how it is laid out."""
    data = [v for v in variables if v[0] not in ("", "None", "__function_workspace__") and not v[0].lower().startswith(("t", "time", "__"))
            and v[2] not in ("char", "object", "opaque", "cell", "struct") and v[1]]
    floats = [v for v in data if v[2] in ("double", "single", "float64", "float32")]
    data = floats or data
    if not data:
        return None
    name, shape, dtype = max(data, key=lambda v: int(np.prod(v[1])) if v[1] else 0)
    if len(shape) == 2 and min(shape) > 1 and min(shape) <= 64:
        return {"variable": name, "layout": "matrix", "n_channels": int(min(shape)), "n_samples": int(max(shape)), "dtype": dtype}
    n = int(np.prod(shape))
    return {"variable": name, "layout": "flat", "n_channels": None, "n_samples": None, "total": n, "dtype": dtype}


def _scan_raw(roots, conn, channels_root=None, **kw):
    cands = []
    for p in _files(roots, (".mat", ".csv", ".h5")):
        base = os.path.basename(p)
        stem, ext = os.path.splitext(base)
        m = _FS_SUFFIX.search(stem)
        facts = {"basename": base, "stem": stem, "ext": ext.lower(), "bytes": os.path.getsize(p), "fs": float(m.group(1)) if m else None,
                 "variables": [], "format": None, "layout": None, "held_out": base == HELD_OUT_FILE, "channels_root": channels_root}
        warnings = []
        try:
            if ext.lower() == ".csv":
                with open(p, encoding="utf-8", errors="replace") as f:
                    first = f.readline()
                ncol = len(first.split(","))
                facts["format"] = "csv"
                facts["variables"] = [{"name": "columns", "shape": [None, ncol], "dtype": "text"}]
                facts["layout"] = {"variable": "columns", "layout": "columns" if ncol > 1 else "flat", "n_channels": ncol if ncol > 1 else None, "n_samples": None}
            else:
                vs, fmt = _mat_variables(p)
                facts["format"] = f"MATLAB {fmt}"
                facts["variables"] = [{"name": n, "shape": s, "dtype": t} for n, s, t in vs]
                facts["layout"] = _layout_guess(vs)
        except Exception as e:
            warnings.append(f"header unreadable: {type(e).__name__}: {e}")
        if facts["fs"] is None:
            warnings.append("fs is not in the file name (_fs<N>): supply it at registration")
        if facts["layout"] and facts["layout"].get("layout") == "flat":
            warnings.append("a flat vector: the channel count must be supplied (this project's convention is 16 end-to-end)")
        facts["derived_dir"] = _derived_dir_for(facts, channels_root)
        cands.append(Candidate(kind="raw", path=p, name=stem, facts=facts, warnings=warnings))
    return cands


def _derived_dir_for(facts, channels_root):
    root = channels_root or KINDS["recording"].roots[0] if "recording" in KINDS else channels_root
    if not root:
        return None
    names = [facts["stem"]] + [v["name"] for v in facts.get("variables", [])]
    for n in names:
        d = os.path.join(root, n)
        if os.path.isdir(d) and _channel_files(d):
            return portable(d)
    for d in glob.glob(os.path.join(root, facts["stem"] + "_fs*")):
        if os.path.isdir(d) and _channel_files(d):
            return portable(d)
    for d in glob.glob(os.path.join(root, "*")):
        man = _read_manifest(d) if os.path.isdir(d) else None
        if man and man.get("source_file") == facts["basename"]:
            return portable(d)
    return None


def _raw_registered_ids(conn, cand: Candidate, channels_root=None, **kw) -> list:
    ids = [r["id"] for r in _recording_rows(conn, cand.facts["basename"])]
    d = cand.facts.get("derived_dir") or _derived_dir_for(cand.facts, channels_root)
    if d:
        ids += [i for i in _recording_ids_for_dir(conn, d) if i not in ids]
    return sorted(ids)


def _check_raw(cand: Candidate, conn, rep: Report, ov: dict, channels_root=None, **kw):
    f = rep.facts
    rep.add("readable", bool(f.get("format")), f.get("format") or "; ".join(cand.warnings))
    rep.add("held_out", not f.get("held_out"), f"{HELD_OUT_FILE} is held out (spec §0 D6): never derived through the interface" if f.get("held_out") else "not the held-out file")
    d = f.get("derived_dir") or _derived_dir_for(f, channels_root or f.get("channels_root"))
    f["derived_dir"] = d
    rep.add("not_derived", not d, f"already derived at {d} — register that recording (kind 'recording') instead of deriving a second copy" if d else "no channel directory yet")
    fs = ov.get("fs", f.get("fs"))
    rep.add("fs", fs is not None and float(fs) > 0 if fs is not None else False, f"fs {fs} Hz" if fs else "fs unknown — supply it")
    if fs is not None:
        f["fs"] = float(fs); f["fs_source"] = ov.get("fs_source", "inferred" if "fs" in ov else "read")
    lay = dict(f.get("layout") or {})
    if ov.get("variable"):
        v = next((x for x in f.get("variables", []) if x["name"] == ov["variable"]), None)
        if v:
            lay = _layout_guess([(v["name"], v["shape"], v["dtype"])]) or lay
    if ov.get("n_channels"):
        lay["n_channels"] = int(ov["n_channels"])
        if lay.get("layout") == "flat" and lay.get("total"):
            lay["n_samples"] = lay["total"] // lay["n_channels"]
    ok = bool(lay.get("variable")) and bool(lay.get("n_channels"))
    rep.add("layout", ok, f"{lay.get('variable')} · {lay.get('layout')} · {lay.get('n_channels')} channels × {lay.get('n_samples') or '?'} samples" if ok
            else "cannot tell the channel layout: supply variable and/or n_channels")
    f["layout"] = lay
    if lay.get("layout") == "flat" and lay.get("total") and lay.get("n_channels") and lay["total"] % lay["n_channels"]:
        rep.add("shape", False, f"{lay['total']:,} samples do not divide into {lay['n_channels']} channels")
    f["source_file"] = f["basename"]
    if ov.get("units") is not None and str(ov.get("units")).strip():
        u = parse_units(ov["units"])
        rep.add("units", u is not None, f"units {u} (supplied)" if u else f"units {ov['units']!r} not recognised — one of {', '.join(UNITS)}")
        f["units"] = u
    else:
        f["units"] = None
        rep.warn(UNITS_UNDECLARED.replace("manifest.json declares", "the raw file declares"))


def _load_raw_channels(path: str, lay: dict):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".csv":
        import pandas as pd
        df = pd.read_csv(path, header=None)
        if df.shape[1] > 1:
            return [np.ascontiguousarray(df.iloc[:, i].to_numpy(dtype=np.float64)) for i in range(df.shape[1])]
        vec = df.iloc[:, 0].to_numpy(dtype=np.float64)
        L = len(vec) // lay["n_channels"]
        return [vec[i * L:(i + 1) * L] for i in range(lay["n_channels"])]
    import scipy.io
    try:
        arr = np.asarray(scipy.io.loadmat(path, variable_names=[lay["variable"]])[lay["variable"]])
    except NotImplementedError:
        import h5py
        with h5py.File(path, "r") as fh:
            arr = np.asarray(fh[lay["variable"]][()])
    if lay["layout"] == "matrix":
        axis = int(np.argmin(arr.shape))
        return [np.ascontiguousarray(np.take(arr, i, axis=axis), dtype=np.float64) for i in range(arr.shape[axis])]
    vec = arr.ravel().astype(np.float64)
    L = len(vec) // lay["n_channels"]
    return [vec[i * L:(i + 1) * L] for i in range(lay["n_channels"])]


def derive_channels(path: str, lay: dict, fs: float, channels_root: str, stem: str | None = None, extra: dict | None = None) -> str:
    """Write ``<channels_root>/<stem>/CH<i>.npy`` + ``manifest.json`` through a
    staging directory swapped in atomically (the layout of
    ``Pipelines/materialize_channels``). Returns the directory."""
    stem = stem or os.path.splitext(os.path.basename(path))[0]
    chans = _load_raw_channels(path, lay)
    lengths = {c.shape[0] for c in chans}
    if len(lengths) != 1:
        raise ValueError(f"channels have inconsistent lengths: {sorted(lengths)}")
    L = lengths.pop()
    os.makedirs(channels_root, exist_ok=True)
    final = os.path.join(channels_root, stem)
    staging = os.path.join(channels_root, f".staging_{stem}_{os.getpid()}")
    if os.path.isdir(staging):
        shutil.rmtree(staging)
    os.makedirs(staging)
    try:
        for i, c in enumerate(chans):
            np.save(os.path.join(staging, f"CH{i}.npy"), c)
        man = {"source_file": os.path.basename(path), "n_channels": len(chans), "fs": fs, "n_samples_per_channel": int(L), "dtype": str(chans[0].dtype),
               "imported_at": now_iso(), "raw_file": portable(path), "variable": lay.get("variable"), "layout": lay.get("layout"),
               "derived_by": "Working.registration (kind raw)", "units": None}
        man.update(extra or {})
        with open(os.path.join(staging, "manifest.json"), "w", encoding="utf-8") as fh:
            json.dump(man, fh, indent=2)
        if os.path.isdir(final):
            shutil.rmtree(final)
        os.replace(staging, final)
    except Exception:
        if os.path.isdir(staging):
            shutil.rmtree(staging)
        raise
    return portable(final)


def _register_raw(conn, cand: Candidate, rep: Report, provenance, actor, writer, channels_root=None, **kw):
    from .core import check as _check
    f = rep.facts
    root = channels_root or f.get("channels_root") or KINDS["recording"].roots[0]
    extra = {}
    if f.get("fs_source") == "inferred":
        extra["fs_note"] = f"{f['fs']} Hz was supplied at registration, not read from the file"
    extra["units"] = f.get("units")
    d = derive_channels(cand.path, f["layout"], float(f["fs"]), root, extra=extra)
    rc = _scan_recording([root], conn)
    rc = next(c for c in rc if _same_path(c.path, d))
    rrep = _check(rc, conn, overrides={"source_file": f["basename"], **({"fs": f["fs"], "fs_source": f["fs_source"]} if f.get("fs_source") == "inferred" else {}),
                                       **({"units": f["units"]} if f.get("units") else {})})
    if not rrep.ok:
        from .core import RegistrationError
        raise RegistrationError("derived channels failed the recording checks: " + "; ".join(f"{c.name}: {c.detail}" for c in rrep.failures()), rrep)
    rid, row = _register_recording(conn, rc, rrep, provenance, actor, writer)
    f["derived_dir"] = d
    f["recording_ids"] = rrep.facts.get("registered_ids")
    f["parameters"] = {"layout": f["layout"], "derived_dir": d}
    # the recording directory gets its own sidecar too
    from .core import build_sidecar
    with open(sidecar_path(d), "w", encoding="utf-8") as fh:
        json.dump(build_sidecar(rc, rrep, {**(provenance or {}), "raw_file": cand.path}, actor, row), fh, indent=2, default=str)
    return rid, row


def _list_raw(conn, sidecar_root=None, **kw):
    return []   # a raw file is never listed as registered on its own: its recording is


# ==================================================================== model ==

def _scan_model(roots, conn, **kw):
    cands = []
    for p in _files(roots, (".pth", ".pt", ".joblib", ".pkl")):
        base = os.path.basename(p)
        ext = os.path.splitext(base)[1].lower()
        fmt = "pytorch" if ext in (".pth", ".pt") else "joblib"
        st = os.stat(p)
        facts = {"format": fmt, "bytes": st.st_size, "mtime": st.st_mtime, "ext": ext}
        m = _HEX8.search(base)
        if m:
            facts["config_hash"] = m.group(1)
        cands.append(Candidate(kind="model", path=p, name=base, facts=facts))
    return cands


def _torch_summary(obj) -> dict:
    import torch
    out = {"type": type(obj).__name__}
    if isinstance(obj, dict):
        out["keys"] = [str(k) for k in list(obj.keys())[:40]]
        for k in ("epoch", "best_val_acc", "val_acc", "accuracy", "arch", "model_name"):
            if k in obj and not hasattr(obj[k], "shape"):
                v = obj[k]
                out[k] = v if isinstance(v, (int, float, str, bool)) else str(v)[:80]
        sd = obj.get("state_dict") if isinstance(obj.get("state_dict"), dict) else obj
        tensors = [v for v in sd.values() if hasattr(v, "numel")]
        out["n_tensors"] = len(tensors)
        out["n_params"] = int(sum(v.numel() for v in tensors))
        if "state_dict" in obj and isinstance(obj["state_dict"], dict):
            out["top_level"] = out["keys"]
            out["keys"] = [str(k) for k in list(obj["state_dict"].keys())[:40]]
    elif isinstance(obj, torch.nn.Module):
        out["n_params"] = int(sum(p.numel() for p in obj.parameters()))
        out["keys"] = list(obj.state_dict().keys())[:40]
    return out


def _sklearn_summary(obj) -> dict:
    out = {"class": type(obj).__name__, "module": type(obj).__module__}
    for k in ("n_features_in_", "n_estimators", "n_classes_", "n_outputs_"):
        if hasattr(obj, k):
            v = getattr(obj, k)
            out[k] = int(v) if isinstance(v, (int, np.integer)) else str(v)
    if hasattr(obj, "classes_"):
        out["classes"] = [c.item() if isinstance(c, np.generic) else c for c in list(getattr(obj, "classes_"))[:20]]
    if hasattr(obj, "feature_names_in_"):
        out["feature_names"] = [str(x) for x in list(getattr(obj, "feature_names_in_"))[:40]]
    if isinstance(obj, dict):
        out["keys"] = [str(k) for k in list(obj.keys())[:40]]
    return out


def _check_model(cand: Candidate, conn, rep: Report, ov: dict, **kw):
    f = rep.facts
    try:
        if f["format"] == "pytorch":
            import torch
            obj = torch.load(cand.path, map_location="cpu", weights_only=False)
            f["summary"] = _torch_summary(obj)
        else:
            import joblib
            obj = joblib.load(cand.path)
            f["summary"] = _sklearn_summary(obj)
        rep.add("readable", True, f"{f['format']} loads: {json.dumps({k: v for k, v in f['summary'].items() if k != 'keys'}, default=str)[:200]}")
    except Exception as e:
        rep.add("readable", False, f"{f['format']} load failed: {type(e).__name__}: {e}")
        return
    rep.add("format", True, f["format"])
    f["parameters"] = {"format": f["format"], "bytes": f["bytes"], "summary": f["summary"], "config_hash": f.get("config_hash")}
    if not os.path.isfile(sidecar_path(cand.path)):
        rep.warn("no provenance sidecar yet: training recipe, data and split are unknown until recorded here")


# ============================================================ matrix profile ==

def _scan_matrix_profile(roots, conn, **kw):
    cands = []
    for root in roots:
        if not os.path.isdir(root):
            continue
        legacy = len(_files([os.path.join(root, "_legacy")], (".npz",)))
        for p in _files([root], (".npz",)):
            base = os.path.basename(p)
            m = _MP_NAME.match(base)
            facts = {"bytes": os.path.getsize(p), "legacy_siblings": legacy}
            warnings = []
            if m:
                facts.update({"stem": m.group("stem"), "channel": int(m.group("ch")), "window_min": float(m.group("win")),
                              "span": [int(m.group("a")), int(m.group("b"))] if m.group("a") else None})
            else:
                warnings.append("file name does not follow mp_v2_<stem>_CH<n>_WIN<len>min[_span<a>-<b>].npz")
            try:
                z = np.load(p, allow_pickle=False)
                mp_shape = _npz_member_shape(p, "mp")
                n_s, mm = _scalar(z, "n_samples"), _scalar(z, "m")
                facts.update({"n_samples": int(n_s) if n_s is not None else None, "m": int(mm) if mm is not None else None,
                              "mp_len": int(mp_shape[0]) if mp_shape else None, "source_file": _scalar(z, "source_file"), "config_hash": _scalar(z, "config_hash")})
                if mp_shape and n_s is not None and mm is not None:
                    span = facts.get("span") or [0, int(n_s)]
                    expect = span[1] - span[0] - int(mm) + 1
                    if int(mp_shape[0]) != expect:
                        warnings.append(f"mp holds {int(mp_shape[0]):,} points; a complete profile over this span has {expect:,} (the length check will fail)")
            except Exception as e:
                warnings.append(f"npz header unreadable: {type(e).__name__}: {e}")
            cands.append(Candidate(kind="matrix_profile", path=p, name=base, facts=facts, warnings=warnings))
    return cands


def _check_matrix_profile(cand: Candidate, conn, rep: Report, ov: dict, **kw):
    f = rep.facts
    try:
        z = np.load(cand.path, allow_pickle=False)
        rep.add("readable", True, f"npz with {len(z.files)} arrays")
    except Exception as e:
        rep.add("readable", False, f"{type(e).__name__}: {e}"); return
    if not _npz_keys(z, rep, ("mp", "m", "fs", "n_samples", "source_file", "channel")):
        return
    sf, ch = _scalar(z, "source_file"), int(_scalar(z, "channel"))
    row = _bind_recording(conn, rep, sf, ch, f.get("stem"))
    stored_id = _scalar(z, "recording_id")
    if row is None:
        return
    if stored_id is not None and int(stored_id) != row["id"]:
        rep.warn(f"the file stores recording_id {stored_id}; bound by content to local id {row['id']} instead")
    fs = _scalar(z, "fs")
    _fs_matches(rep, fs, row)
    f["fs"] = float(fs)
    n_file = int(_scalar(z, "n_samples"))
    if n_file != int(row["n_samples"]):
        rep.warn(f"the file says n_samples {n_file:,}; the recording has {row['n_samples']:,}")
    span = f.get("span") or ([int(_scalar(z, "span_start")), int(_scalar(z, "span_end"))] if "span_start" in z.files else [0, int(row["n_samples"])])
    _span_inside(rep, span, row)
    m = int(_scalar(z, "m"))
    mp = z["mp"]
    expect = (span[1] - span[0]) - m + 1
    rep.add("length", len(mp) == expect, f"len(mp) {len(mp):,} == span {span[1] - span[0]:,} − m {m} + 1" if len(mp) == expect
            else f"len(mp) {len(mp):,} ≠ span {span[1] - span[0]:,} − m {m} + 1 = {expect:,}")
    _finite(mp, rep, "mp")
    f["parameters"] = {"m": m, "window_min": _scalar(z, "window_min", f.get("window_min")), "config_hash": _scalar(z, "config_hash"),
                       "backend": _scalar(z, "backend"), "approx": _scalar(z, "approx"), "created_at": _scalar(z, "created_at"),
                       "data_sha1": _scalar(z, "data_sha1"), "elapsed_s": _scalar(z, "elapsed_s"), "n_samples_in_file": n_file}
    f["config_hash"] = _scalar(z, "config_hash")
    f["producer"] = "Adapters/detection_matrix_profile.py (mp_v2)"


# ============================================================ window matrix ==

def _scan_window_matrix(roots, conn, **kw):
    cands = []
    for p in _files(roots, (".npz", ".csv")):
        base = os.path.basename(p)
        facts = {"bytes": os.path.getsize(p), "format": "npz" if base.endswith(".npz") else "csv"}
        warnings = []
        m = _WM_NAME.match(base)
        if m:
            facts.update({"stem": m.group("stem"), "channel": int(m.group("ch")), "window_min": float(m.group("win")), "step_pct": float(m.group("step"))})
        elif facts["format"] == "npz":
            warnings.append("file name does not follow wm_v1_<stem>_CH<n>_WIN<len>min_STEP<pct>pct.npz")
        else:
            warnings.append("legacy CSV window matrix: no manifest fields; the recording is inferred from the file name and must be checked by eye")
        cands.append(Candidate(kind="window_matrix", path=p, name=base, facts=facts, warnings=warnings))
    return cands


def _check_window_matrix(cand: Candidate, conn, rep: Report, ov: dict, **kw):
    f = rep.facts
    if f["format"] == "csv":
        try:
            with open(cand.path, encoding="utf-8", errors="replace", newline="") as fh:
                r = csv.reader(fh)
                header = next(r)
                n = sum(1 for _ in r)
            rep.add("readable", True, f"csv: {len(header)} columns, {n:,} rows")
        except Exception as e:
            rep.add("readable", False, f"{type(e).__name__}: {e}"); return
        f["columns"], f["n_rows"] = header, n
        f["parameters"] = {"columns": header[:64], "n_rows": n, "format": "csv"}
        stems = []
        if conn is not None:
            for sf in {r["source_file"] for r in _recording_rows(conn)}:
                st = sf.rsplit(".", 1)[0]
                if st and st in cand.name:
                    stems.append(sf)
        if stems:
            row = _recording_rows(conn, stems[0])[0]
            rep.add("recording", True, f"file name mentions {stems[0]} (recordings id {row['id']} CH{row['channel']}); channel not encoded in the name")
            f["recording_id"] = row["id"]; f["source_file"] = stems[0]
        else:
            rep.warn("no registered recording is named in the file name")
        return
    try:
        z = np.load(cand.path, allow_pickle=False)
        rep.add("readable", True, f"npz with {len(z.files)} arrays")
    except Exception as e:
        rep.add("readable", False, f"{type(e).__name__}: {e}"); return
    if not _npz_keys(z, rep, ("values", "computed", "start_idx", "m", "fs", "source_file", "channel", "n_samples")):
        return
    values, computed, starts = z["values"], z["computed"], z["start_idx"]
    ok = values.ndim == 2 and computed.shape == values.shape and len(starts) == values.shape[0]
    rep.add("shape", ok, f"values {values.shape}, computed {computed.shape}, {len(starts)} starts" + ("" if ok else " — disagree"))
    row = _bind_recording(conn, rep, _scalar(z, "source_file"), int(_scalar(z, "channel")), f.get("stem"))
    if row is None:
        return
    fs = _scalar(z, "fs"); _fs_matches(rep, fs, row); f["fs"] = float(fs)
    m = int(_scalar(z, "m"))
    span = [int(_scalar(z, "span_start", 0)), int(_scalar(z, "span_end", row["n_samples"]))]
    span[1] = max(span[1], int(starts.max()) + m if len(starts) else span[1])
    _span_inside(rep, span, row)
    if len(starts):
        _finite(values[computed], rep, "computed values") if computed.any() else rep.add("finite", True, "no computed cells")
    else:
        rep.add("finite", True, "no windows")
    complete = bool(_scalar(z, "complete", True)) and bool(computed.all()) if computed.size else False
    if not complete:
        rep.warn(f"incomplete: {int(computed.sum()):,} of {computed.size:,} cells computed" if computed.size else "no cells computed")
    f["n_windows"] = int(values.shape[0])
    f["parameters"] = {"m": m, "step": _scalar(z, "step"), "window_min": _scalar(z, "window_min"), "step_frac": _scalar(z, "step_frac"), "n_windows": f["n_windows"],
                       "n_columns": int(values.shape[1]), "columns": [str(c) for c in z["columns"][:64]] if "columns" in z.files else None,
                       "complete": complete, "config_hash": _scalar(z, "config_hash"), "created_at": _scalar(z, "created_at"), "builder_version": _scalar(z, "builder_version")}
    f["config_hash"] = _scalar(z, "config_hash")
    f["producer"] = "Adapters/preprocessing_window_matrix.py (wm_v1)"


# ============================================================== window set ==

def _scan_dirs_with(roots, filename, kind, max_depth=3):
    cands = []
    for root in roots:
        if not os.path.isdir(root):
            continue
        base_depth = root.rstrip("/\\").count(os.sep) + root.rstrip("/\\").count("/")
        for r, dirs, files in os.walk(root):
            dirs[:] = sorted(d for d in dirs if not d.startswith("."))
            depth = r.count(os.sep) + r.count("/") - base_depth
            if depth > max_depth:
                dirs[:] = []
                continue
            if filename in files:
                cands.append(Candidate(kind=kind, path=portable(r), name=os.path.basename(r), facts={"files": sorted(files)}))
                dirs[:] = []
    return cands


def _scan_window_set(roots, conn, **kw):
    out = []
    for c in _scan_dirs_with(roots, "manifest.json", "window_set"):
        man = _read_manifest(c.path) or {}
        c.facts.update({"manifest": man, "n_windows": man.get("n_windows"), "length": man.get("length"), "fs": man.get("fs"), "source_file": man.get("recording"), "channel": man.get("channel")})
        if man.get("error"):
            c.warnings.append(f"manifest.json unreadable: {man['error']}")
        out.append(c)
    return out


def _check_window_set(cand: Candidate, conn, rep: Report, ov: dict, **kw):
    f = rep.facts
    man = f.get("manifest") or {}
    if not man or man.get("error"):
        rep.add("manifest", False, man.get("error") or "no manifest.json"); return
    rep.add("manifest", man.get("kind") == "window_set", f"kind {man.get('kind')!r}" + ("" if man.get("kind") == "window_set" else " — expected 'window_set'"))
    p = os.path.join(cand.path, "windows.npz")
    if not rep.add("files", os.path.isfile(p), "windows.npz present" if os.path.isfile(p) else "windows.npz missing"):
        return
    try:
        z = np.load(p, allow_pickle=False)
        rep.add("readable", True, f"windows.npz with {len(z.files)} arrays")
    except Exception as e:
        rep.add("readable", False, f"{type(e).__name__}: {e}"); return
    if not _npz_keys(z, rep, ("starts", "length", "fs", "source_file", "channel")):
        return
    starts, length = z["starts"], int(_scalar(z, "length"))
    row = _bind_recording(conn, rep, _scalar(z, "source_file"), int(_scalar(z, "channel")))
    if row is None:
        return
    fs = _scalar(z, "fs"); _fs_matches(rep, fs, row); f["fs"] = float(fs)
    n_w = int(len(starts))
    rep.add("shape", man.get("n_windows") in (None, n_w), f"{n_w} windows × {length} samples" if man.get("n_windows") in (None, n_w) else f"manifest says {man.get('n_windows')} windows, npz holds {n_w}")
    _span_inside(rep, [int(starts.min()) if n_w else 0, (int(starts.max()) + length) if n_w else length], row)
    f["n_windows"] = n_w
    f["parameters"] = {"length": length, "n_windows": n_w, "labels_source": man.get("labels_source"), "has_labels": "labels" in z.files, "name": man.get("name")}


# ================================================================ encoding ==

def _scan_encoding(roots, conn, **kw):
    cands = []
    for p in _files(roots, (".npz",)):
        cands.append(Candidate(kind="encoding", path=p, name=os.path.basename(p), facts={"bytes": os.path.getsize(p)}))
    return cands


def _encoding_registered_ids(conn, cand: Candidate, **kw):
    return [r["id"] for r in _rows(conn, "SELECT id, path FROM encodings WHERE active = 1") if _same_path(r["path"], cand.path)]


def _check_encoding(cand: Candidate, conn, rep: Report, ov: dict, **kw):
    f = rep.facts
    try:
        z = np.load(cand.path, allow_pickle=False)
        rep.add("readable", True, f"npz with {len(z.files)} arrays")
    except Exception as e:
        rep.add("readable", False, f"{type(e).__name__}: {e}"); return
    if not _npz_keys(z, rep, ("values", "encoding_type", "source_file", "channel", "fs", "span_start", "span_end", "config_hash")):
        return
    row = _bind_recording(conn, rep, _scalar(z, "source_file"), int(_scalar(z, "channel")))
    if row is None:
        return
    fs = _scalar(z, "fs"); _fs_matches(rep, fs, row); f["fs"] = float(fs)
    _span_inside(rep, [int(_scalar(z, "span_start")), int(_scalar(z, "span_end"))], row)
    _finite(z["values"], rep, "values")
    f["encoding_type"] = str(_scalar(z, "encoding_type")); f["config_hash"] = str(_scalar(z, "config_hash"))
    f["parameters"] = {"encoding_type": f["encoding_type"], "config_hash": f["config_hash"], "shape": list(z["values"].shape), "dtype": str(z["values"].dtype)}


def _register_encoding(conn, cand: Candidate, rep: Report, provenance, actor, writer, **kw):
    f = rep.facts
    row = {"recording_id": f["recording_id"], "span_start": f["span"][0], "span_end": f["span"][1], "encoding_type": f["encoding_type"],
           "config_hash": f["config_hash"], "path": cand.path, "created_at": now_iso(), "active": 1}
    return writer(conn, "encodings", row), row


def _unregister_encoding(conn, row_id, actor):
    conn.execute("UPDATE encodings SET active = 0 WHERE id = ?", (int(row_id),))
    return {"table": "encodings", "id": int(row_id), "active": 0}


def _list_encodings(conn, sidecar_root=None, **kw):
    out = []
    for r in _rows(conn, "SELECT * FROM encodings WHERE active = 1 ORDER BY id"):
        r.update({"kind": "encoding", "name": os.path.basename(r["path"]), "exists": os.path.isfile(r["path"]), "manifest": read_sidecar(r["path"], sidecar_root)})
        out.append(r)
    return out


# ======================================================== drop-motif store ==

DROP_REQUIRED = ("event_id", "recording_id", "source_file", "channel", "fs")


def _scan_drop_store(roots, conn, **kw):
    out = []
    for c in _scan_dirs_with(roots, "events.csv", "drop_motif_store"):
        man = _read_manifest(c.path) or {}
        c.facts.update({"detector": man.get("detector") or man.get("kind"), "n_motifs": man.get("n_motifs"), "manifest": man})
        out.append(c)
    return out


def _check_drop_store(cand: Candidate, conn, rep: Report, ov: dict, **kw):
    f = rep.facts
    need = {"events.csv", "snippets.npz", "manifest.json"}
    have = set(os.listdir(cand.path))
    rep.add("files", need <= have, f"missing {sorted(need - have)}" if not need <= have else "events.csv + snippets.npz + manifest.json")
    if not need <= have:
        return
    man = _read_manifest(cand.path) or {}
    rep.add("manifest", bool(man) and "error" not in man, man.get("error", f"detector {man.get('detector')!r}"))
    try:
        with open(os.path.join(cand.path, "events.csv"), encoding="utf-8", newline="") as fh:
            rows = list(csv.DictReader(fh))
        rep.add("readable", True, f"{len(rows):,} events")
    except Exception as e:
        rep.add("readable", False, f"events.csv: {type(e).__name__}: {e}"); return
    cols = list(rows[0].keys()) if rows else []
    missing = [c for c in DROP_REQUIRED if c not in cols]
    rep.add("columns", not missing, f"missing columns {missing}" if missing else f"{len(cols)} columns")
    if missing:
        return
    key_col = "snippet_key" if "snippet_key" in cols else "event_id"
    try:
        z = np.load(os.path.join(cand.path, "snippets.npz"), allow_pickle=False)
        keys = set(z.files)
    except Exception as e:
        rep.add("snippets", False, f"snippets.npz: {type(e).__name__}: {e}"); return
    # store.py writes three arrays per event, `<snippet_key>__raw_mv`, `__detrended_mv`, `__t_s`; an older
    # store may hold one array under the bare key — either counts as "has its snippet"
    prefixes = {k.split("__", 1)[0] for k in keys}
    lost = [r[key_col] for r in rows if r[key_col] not in keys and r[key_col] not in prefixes]
    rep.add("snippets", not lost, f"{len(lost)} events without a snippet: {lost[:5]}" if lost else f"every one of {len(rows)} events has its snippet ({len(keys)} arrays in the archive)")
    rec_ids = sorted({int(float(r["recording_id"])) for r in rows if r["recording_id"] not in ("", None)})
    if conn is not None:
        for rid in rec_ids:
            row = conn.execute("SELECT id, fs, source_file FROM recordings WHERE id = ? AND active = 1", (rid,)).fetchone()
            if row is None:
                rep.warn(f"recording_id {rid} named in events.csv is not a registered recording here (another machine's id?)")
            else:
                fs_ev = {float(r["fs"]) for r in rows if int(float(r["recording_id"])) == rid and r["fs"]}
                if fs_ev and any(abs(x - float(row[1])) > 1e-9 for x in fs_ev):
                    rep.warn(f"events on recording_id {rid} say fs {sorted(fs_ev)}; the recording has {row[1]}")
    f.update({"n_events": len(rows), "n_snippets": len(keys), "detector": man.get("detector") or man.get("kind"), "recording_ids": rec_ids,
              "spans": sorted({r.get("span_key") for r in rows if r.get("span_key")}), "columns": cols})
    if len(rec_ids) == 1 and conn is not None:
        r0 = conn.execute("SELECT id FROM recordings WHERE id = ? AND active = 1", (rec_ids[0],)).fetchone()
        if r0:
            f["recording_id"] = rec_ids[0]
    f["parameters"] = {"n_events": len(rows), "n_snippets": len(keys), "detector": f["detector"], "recording_ids": rec_ids, "spans": f["spans"],
                       "manifest_params": man.get("params") or man.get("parameters"), "n_pure": man.get("n_pure"), "n_impure": man.get("n_impure")}
    f["producer"] = f"Working/Detection/drop_motifs ({f['detector']})"


# ===================================================== catalogue spreadsheet ==

CATALOGUE_REQUIRED = ("ID_Number", "ID_Name", "Channel", "StartTime_h", "StopTime_h", "DATASET", "STATUS")
_CAT_INT = ("ID_Number", "Channel", "Pack", "Parent_ID")
_CAT_FLOAT = ("StartTime_h", "StopTime_h", "Length_h", "Length_m", "Length_s", "AmplitudeMax", "AmplitudeMin")


def _scan_catalogue(roots, conn, **kw):
    return [Candidate(kind="catalogue_spreadsheet", path=p, name=os.path.basename(p), facts={"bytes": os.path.getsize(p)}) for p in _files(roots, (".xlsx",))]


def _check_catalogue(cand: Candidate, conn, rep: Report, ov: dict, **kw):
    import openpyxl
    f = rep.facts
    try:
        wb = openpyxl.load_workbook(cand.path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        raw = list(ws.iter_rows(values_only=True))
        rep.add("readable", True, f"sheet {ws.title!r}: {len(raw)} rows")
    except Exception as e:
        rep.add("readable", False, f"{type(e).__name__}: {e}"); return
    if not raw:
        rep.add("columns", False, "empty sheet"); return
    header = [str(h).strip() if h is not None else "" for h in raw[0]]
    missing = [c for c in CATALOGUE_REQUIRED if c not in header]
    rep.add("columns", not missing, f"missing required columns {missing}" if missing else f"{len(header)} columns")
    if missing:
        return
    rows, bad = [], []
    for i, line in enumerate(raw[1:], start=2):
        if all(v is None for v in line):
            continue
        rec = {}
        for col, v in zip(header, line):
            if not col:
                continue
            if v is None:
                rec[col] = None; continue
            try:
                if col in _CAT_INT:
                    if isinstance(v, str) and not re.fullmatch(r"\s*-?\d+\s*", v):
                        raise ValueError(v)
                    rec[col] = int(v)
                elif col in _CAT_FLOAT:
                    if isinstance(v, str) and not re.fullmatch(r"\s*-?\d+(\.\d+)?\s*", v):
                        raise ValueError(v)
                    rec[col] = float(v)
                else:
                    rec[col] = v if isinstance(v, (int, float)) else str(v)
            except (TypeError, ValueError):
                bad.append({"row": i, "column": col, "value": str(v)})
                rec[col] = None
        rows.append(rec)
    if bad:
        rep.warn(f"{len(bad)} unparseable cell{'s' if len(bad) != 1 else ''} left empty, not guessed: " + "; ".join(f"row {b['row']} {b['column']}={b['value']!r}" for b in bad[:8]))
    f.update({"n_rows": len(rows), "columns": header, "unparseable": bad, "rows": rows})
    f["parameters"] = {"n_rows": len(rows), "columns": header, "unparseable": bad, "datasets": sorted({str(r.get("DATASET")) for r in rows if r.get("DATASET")})}


# ============================================================== HPC results ==

_PAUSED_STATUSES = ("paused", "waiting", "pending", "queued", "hpc")


def _scan_hpc(roots, conn, **kw):
    cands = []
    for root in roots:
        if not os.path.isdir(root):
            continue
        for name in sorted(os.listdir(root)):
            d = os.path.join(root, name)
            if not os.path.isdir(d) or name.startswith("."):
                continue
            files = sorted(os.listdir(d))
            jsons = [x for x in files if x.endswith(".json") and x != "manifest.json" and not x.endswith(".manifest.json")]
            if not jsons:
                continue
            jsons.sort(key=lambda x: (x != f"{name}.json", x))   # <job>.json first when there are several
            facts = {"files": files, "recipe_file": jsons[0], "artifacts": [x for x in files if x.endswith((".npz", ".npy", ".csv", ".joblib", ".pth"))],
                     "has_manifest": "manifest.json" in files, "recipe_hash": None}
            warnings = []
            try:
                with open(os.path.join(d, jsons[0]), encoding="utf-8") as fh:
                    recipe = json.load(fh)
                from Working.recipes import short_hash
                facts["recipe"] = recipe
                facts["recipe_hash"] = short_hash(recipe)
                facts["recording_id_in_recipe"] = recipe.get("recording_id")
            except Exception as e:
                warnings.append(f"{jsons[0]}: {type(e).__name__}: {e}")
            cands.append(Candidate(kind="hpc_result", path=portable(d), name=name, facts=facts, warnings=warnings))
    return cands


def _check_hpc(cand: Candidate, conn, rep: Report, ov: dict, **kw):
    f = rep.facts
    recipe = f.get("recipe")
    ok = isinstance(recipe, dict) and isinstance(recipe.get("steps"), list) and bool(recipe.get("steps"))
    rep.add("recipe", ok, f"{f.get('recipe_file')}: {len(recipe['steps'])} step(s)" if ok else f"{f.get('recipe_file')} is not a recipe: " + "; ".join(cand.warnings))
    if not ok:
        return
    h = f.get("recipe_hash")
    rep.add("recipe_hash", bool(h), h or "recipe hash could not be computed")
    named = _HEX8.search(cand.name)
    if h and named and named.group(1) != h:
        rep.warn(f"directory name carries {named.group(1)} but the recipe hashes to {h}")
    # the recording: by CONTENT when an artifact names it (source_file + channel), else the recipe's local
    # recording_id (this database's id — a recipe exported elsewhere needs overrides["recording_id"])
    row = None
    rid = ov.get("recording_id", recipe.get("recording_id"))
    content = None
    for a in f.get("artifacts", []):
        if a.endswith(".npz"):
            try:
                z = np.load(os.path.join(cand.path, a), allow_pickle=False)
                if "source_file" in z.files and "channel" in z.files:
                    content = (_scalar(z, "source_file"), int(_scalar(z, "channel"))); break
            except Exception:
                continue
    if conn is not None and content is not None:
        rows = _recording_rows(conn, content[0], content[1])
        row = rows[0] if rows else None
        rep.add("recording", row is not None, f"bound by content to {content[0]} CH{content[1]}" + (f" = recordings id {row['id']}" if row else " — no such registered recording"))
        if row is not None and rid is not None and int(rid) != row["id"]:
            rep.warn(f"the recipe says recording_id {rid}; the artifacts name {content[0]} CH{content[1]} (local id {row['id']}) — bound by content")
    else:
        if conn is not None and rid is not None:
            rows = _rows(conn, "SELECT * FROM recordings WHERE id = ? AND active = 1", (int(rid),))
            row = rows[0] if rows else None
        rep.add("recording", row is not None, f"recording_id {rid} = {row['source_file']} CH{row['channel']}" if row
                else f"recording_id {rid} is not a registered recording here (ids are local: pass overrides.recording_id, or let an artifact name source_file + channel)")
    if row:
        f["recording_id"] = row["id"]; f["source_file"] = row["source_file"]; f["channel"] = row["channel"]; f["fs"] = row["fs"]
    arts = []
    shape_ok = length_ok = finite_ok = True
    details = []
    for a in f.get("artifacts", []):
        p = os.path.join(cand.path, a)
        info = {"name": a, "bytes": os.path.getsize(p)}
        if a.endswith(".npz"):
            try:
                z = np.load(p, allow_pickle=False)
            except Exception as e:
                details.append(f"{a}: {type(e).__name__}: {e}"); shape_ok = False; arts.append(info); continue
            n_file = _scalar(z, "n_samples")
            if row is not None and n_file is not None and int(n_file) != int(row["n_samples"]):
                shape_ok = False; details.append(f"{a}: n_samples {int(n_file):,} ≠ recording {row['n_samples']:,}")
            main = next((k for k in ("mp", "values", "scores", "labels") if k in z.files), None)
            if main is not None:
                arr = z[main]
                m = _scalar(z, "m")
                if row is not None and main == "mp" and m is not None:
                    span = [int(_scalar(z, "span_start", 0)), int(_scalar(z, "span_end", row["n_samples"]))]
                    expect = span[1] - span[0] - int(m) + 1
                    if len(arr) != expect:
                        length_ok = False; details.append(f"{a}: len(mp) {len(arr):,} ≠ {expect:,}")
                if np.asarray(arr).dtype.kind in "fc" and np.isnan(arr).any():
                    finite_ok = False; details.append(f"{a}: NaN in {main}")
                info.update({"main": main, "shape": list(np.asarray(arr).shape), "n_samples": int(n_file) if n_file is not None else None, "config_hash": _scalar(z, "config_hash")})
        arts.append(info)
    rep.add("shape", shape_ok, "; ".join(d for d in details if "n_samples" in d or "Error" in d) or "per-channel shapes match the recording")
    rep.add("length", length_ok, "; ".join(d for d in details if "len(" in d) or "lengths agree with m and the span")
    rep.add("finite", finite_ok, "; ".join(d for d in details if "NaN" in d) or "no NaN in the artifacts")
    f["artifacts"] = arts
    if f.get("has_manifest"):
        man = _read_manifest(cand.path) or {}
        ok_m = "error" not in man
        rep.add("manifest", ok_m, man.get("error") or (f"manifest.json with {len(man['runs'])} run(s)" if isinstance(man.get("runs"), list) else "manifest.json (not a run manifest; kept as provenance only)"))
        if ok_m and not isinstance(man.get("runs"), list):
            rep.warn("manifest.json is not a Pipelines/run_recipe manifest (no 'runs' list): nothing is imported from it")
            f["has_manifest"] = False
        elif ok_m and h and any(r.get("config_hash") not in (None, h) for r in man["runs"]):
            rep.warn("manifest.json config_hash differs from the recipe hash")
    f["paused_run_id"] = None
    if conn is not None and h:
        placeholders = ",".join("?" for _ in _PAUSED_STATUSES)
        r = conn.execute(f"SELECT r.id FROM runs r JOIN configs c ON c.id = r.config_id WHERE c.config_hash = ? AND r.status IN ({placeholders}) ORDER BY r.id DESC LIMIT 1",
                         (h, *_PAUSED_STATUSES)).fetchone()
        if r is not None:
            f["paused_run_id"] = int(r[0])
            rep.warn(f"recipe hash {h} matches paused run {r[0]}: the checks Jobs › Upload shows pass here; continuing the run is the job model's step")
    f["parameters"] = {"recipe_hash": h, "recipe_file": f.get("recipe_file"), "steps": [f"{s.get('stage')}.{s.get('algorithm')}" for s in recipe["steps"]],
                       "artifacts": arts, "paused_run_id": f["paused_run_id"], "has_manifest": f.get("has_manifest")}
    f["producer"] = f"HPC job {cand.name}"


def _register_hpc(conn, cand: Candidate, rep: Report, provenance, actor, writer, **kw):
    """The bundle row, plus one row per artifact inside it under its OWN kind
    (an mp_v2_*.npz becomes a registered matrix_profile in place), so the
    consumers that read the registry see it without a copy."""
    from .core import check as _check
    f = rep.facts
    if f.get("has_manifest"):
        try:
            from Working.manifest import import_manifest
            f["parameters"]["imported"] = import_manifest(conn, os.path.join(cand.path, "manifest.json"))
        except Exception as e:
            f["parameters"]["imported"] = {"error": f"{type(e).__name__}: {e}"}
    published = []
    for a in f.get("artifacts", []):
        name = a["name"]
        kind = "matrix_profile" if _MP_NAME.match(name) else "window_matrix" if _WM_NAME.match(name) else None
        if kind is None:
            continue
        sub = Candidate(kind=kind, path=portable(os.path.join(cand.path, name)), name=name)
        for c in KINDS[kind].scan([cand.path], conn):
            if c.name == name:
                sub = c
        srep = _check(sub, conn)
        if srep.ok:
            sid, _row = _art_register(conn, sub, srep, {**(provenance or {}), "producer": f"HPC job {cand.name}"}, actor, writer)
            published.append({"name": name, "kind": kind, "id": sid})
        else:
            published.append({"name": name, "kind": kind, "id": None, "failed": [c.name for c in srep.failures()]})
    f["parameters"]["published"] = published
    return _art_register(conn, cand, rep, provenance, actor, writer)


# ================================================================= registry ==

def _art_kind(name, label, roots, ui, naming, scan, check, register=_art_register):
    return KindSpec(name=name, label=label, roots=roots, table="registered_artifacts", ui=ui, naming=naming, scan=scan, check=check,
                    register=register, registered_ids=_art_registered_ids, unregister=_art_unregister, list_registered=_art_list(name))


KINDS: dict = {}
KINDS["recording"] = KindSpec(
    name="recording", label="Recording", roots=["DATA/derived/channels"], table="recordings",
    ui="Settings › Datasets; Explore recording menu (GET /api/recordings); Analyse source picker; Channels & events",
    naming="DATA/derived/channels/<stem>/CH<n>.npy + manifest.json (source_file, fs, n_channels, n_samples_per_channel, fs_note, time_base, units)",
    scan=_scan_recording, check=_check_recording, register=_register_recording, registered_ids=_recording_registered_ids,
    unregister=_unregister_recording, list_registered=_list_recordings, enrich=_enrich_recording)
KINDS["raw"] = KindSpec(
    name="raw", label="Raw recording file", roots=["DATA/raw"], table="recordings",
    ui="Settings › Datasets › Import a recording (derives channels, then registers the recording)",
    naming="DATA/raw/<stem>[_fs<N>].mat | .csv — flat vector (16 channels end to end), (n_samples × n_channels) matrix, or one column per channel",
    scan=_scan_raw, check=_check_raw, register=_register_raw, registered_ids=_raw_registered_ids,
    unregister=_unregister_recording, list_registered=_list_raw)
KINDS["model"] = _art_kind("model", "Model", ["MODELS", "DATA/derived/models"], "Settings › Models & registration › Registered models; Models registry",
                           "MODELS/<name>.pth (PyTorch checkpoint) · DATA/derived/models/catalogue_classifier_<hash16>.joblib (sklearn)",
                           _scan_model, _check_model)
KINDS["window_matrix"] = _art_kind("window_matrix", "Window matrix", ["Results/Preprocessing/window_matrix", "MATRICES"],
                                   "Settings › Storage (window matrices root); Analyse source; Models › Launch sources",
                                   "wm_v1_<stem>_CH<n>_WIN<len>min_STEP<pct>pct.npz · legacy MATRICES/*.csv", _scan_window_matrix, _check_window_matrix)
KINDS["matrix_profile"] = _art_kind("matrix_profile", "Matrix profile", ["Results/Detection/matrix_profile"],
                                    "Settings › Storage (matrix profiles root); Analyse matrix_profile resume; Discovery",
                                    "mp_v2_<stem>_CH<n>_WIN<len>min[_span<a>-<b>].npz", _scan_matrix_profile, _check_matrix_profile)
KINDS["window_set"] = _art_kind("window_set", "Window set", ["DATA/derived/window_sets"], "Library › Window sets; Models › Launch",
                                "DATA/derived/window_sets/<name>/windows.npz + manifest.json (kind window_set)", _scan_window_set, _check_window_set)
KINDS["encoding"] = KindSpec(
    name="encoding", label="Encoding", roots=["DATA/derived/encodings"], table="encodings", ui="Analyse (Encoding consumers); Settings › Storage",
    naming="DATA/derived/encodings/enc_<type>_<stem>_CH<n>_<hash8>.npz (values, encoding_type, source_file, channel, fs, span_start, span_end, config_hash)",
    scan=_scan_encoding, check=_check_encoding, register=_register_encoding, registered_ids=_encoding_registered_ids,
    unregister=_unregister_encoding, list_registered=_list_encodings)
KINDS["drop_motif_store"] = _art_kind("drop_motif_store", "Drop-motif event store", ["DATA/derived/drop_motifs", "DATA/library_seed", "Plots"],
                                      "Library import (Prompt 03); Settings › Storage", "<dir>/events.csv + snippets.npz + manifest.json (Working/Detection/drop_motifs/store.py)",
                                      _scan_drop_store, _check_drop_store)
KINDS["catalogue_spreadsheet"] = _art_kind("catalogue_spreadsheet", "Catalogue spreadsheet", ["DATA/catalogue"], "Library import (Prompt 03); Settings › Storage",
                                           "DATA/catalogue/*.xlsx with columns ID_Number, ID_Name, Channel, StartTime_h, StopTime_h, DATASET, STATUS, …",
                                           _scan_catalogue, _check_catalogue)
KINDS["hpc_result"] = _art_kind("hpc_result", "HPC result bundle", ["HPC/results"], "Jobs › Upload (Prompt 01 continues the paused run); Settings › Storage",
                                "HPC/results/<job>/<job>.json (the recipe) + artifacts (*.npz) [+ manifest.json from Pipelines/run_recipe]", _scan_hpc, _check_hpc, _register_hpc)
