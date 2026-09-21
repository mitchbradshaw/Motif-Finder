"""
catalogue.py
============
`DATA/catalogue/signal_catalog.xlsx` -> `annotations` + `sequences`.
`docs/LIBRARY_STORAGE.md` §5, the `catalogue` row of the importer table.

The spreadsheet is a researcher's notebook, not an export: 37 rows, 26 columns,
free text in five of them and a foreign key in none of them. Three things in it
bind to the database, and each one is a rule rather than a lookup.

**To a recording.** `Pack` and `Channel` are *per-pack* — `Channel` runs 0..3
inside a pack — so the global channel this project's `recordings` table keys on
is ``Pack * 4 + Channel``. That rule was verified against the drop-motif seed's
own `catalogue_id -> (recording_id, channel)` mapping for all 15 catalogue ids
the seed contains, 15/15. Binding on `Channel` alone would silently put seven
of the catalogue's rows on the wrong trace, which is the failure this module's
`global_channel` exists to make impossible to write by accident.

**To a span.** `StartTime_h * 3600 * fs` -> `start_idx`. `fs` is read off the
`recordings` row rather than hard-coded at 1.0: the catalogue's recordings all
happen to be 1 Hz, but a constant here would be a silent unit bug the day one
is not.

**To a species.** There is no species column. Species is `(Experiment, Pack)`
decoded against the free-text legend in `Reference_Notes`, and the two
experiments give `Pack` *different* meanings — E01 Pack 0 is Lion's mane, E02
Pack 0 is White Oyster. So the legend is hard-coded below as an explicit
mapping with its evidence quoted, keyed on the pair, and a pair that is not in
it gets **no species tag at all**. An unknown mushroom is not a guessable one.

What this importer writes, and what it deliberately does not:

* an `annotations` row (human table) and a `sequences` row with
  `origin = 'human'`, `source_kind = 'catalogue'`. A catalogue row is a
  person's claim; it writes no `detections` row and never will (CLAUDE.md
  rule 5);
* `needs_extraction = 1` on every sequence, because a catalogue row states
  *"16 cycles"* and does not say where the sixteen are. The claim is recorded,
  the events are not invented (§4);
* nothing at all for `M4_aug_concat_fs1.mat`, the held-out evaluation
  recording, which is refused on every route.

**Idempotence** keys on `(recording_id, start_idx, end_idx, source)` in
`annotations` — `Working.database.queries.annotation_exists`. That is the key
the 31 rows already in the live database carry, and they were written before
this module existed, so keying on anything this importer alone produces (an id
column, a hash of the row) would re-import all 31 and double the catalogue.

Nothing here imports a UI library (CLAUDE.md rule 1).
"""

import os
import re

from Working.config import HELD_OUT_RECORDING_FILE
from Working.database import queries as q
from Working.database.vocabulary import SEED_VOCABULARY, get_or_create_term

#: The `annotations.source` value catalogue rows carry. Fixed by the 31 rows
#: already in the live database — changing it would orphan every one of them.
CATALOGUE_SOURCE = "excel_catalog"

#: Channels per pack. The acquisition wires four electrodes into each pack and
#: the spreadsheet numbers them within the pack, so this is a property of the
#: rig, not a tuning parameter.
CHANNELS_PER_PACK = 4

#: `(Experiment, Pack) -> species`, decoded from the `Reference_Notes` legend.
#:
#: E01: "R0 GT01 is Grow Tent; R1 Pack 0 is Lions mane (Urban Valley); Pack 1
#:       is Chocolate Oyster; Pack 2 is Snow White Oyster; Pack 3 is Shitaki"
#: E02: "EO2 is November 2025; Pack 0 is White Oyster (Little Acare); Pack 1
#:       Pink Oyster; Pack 2 Pink Oyster; Pack 3 Golden Oyster;"
#:
#: Note E02 gives Pack 1 and Pack 2 the *same* species — that is what the
#: legend says, not a transcription slip, and it is kept rather than
#: "corrected". "Shitaki" is the notebook's spelling of shiitake; the tag is
#: spelled correctly because the tag is the project's vocabulary, while the
#: legend is the evidence for it.
SPECIES_BY_EXPERIMENT_PACK = {
    ("E01", 0): "lions_mane",
    ("E01", 1): "chocolate_oyster",
    ("E01", 2): "snow_white_oyster",
    ("E01", 3): "shiitake",
    ("E02", 0): "white_oyster",
    ("E02", 1): "pink_oyster",
    ("E02", 2): "pink_oyster",
    ("E02", 3): "golden_oyster",
}

#: `Elements` is a semi-controlled vocabulary typed by hand over two years, so
#: it carries typos. Only corrections with evidence in the sheet go here; a
#: value this map does not know is created as a new term rather than dropped,
#: because a silently discarded label is a lost observation.
ELEMENT_TYPOS = {
    "stegasauras": "stegasaurus",
}

#: Cells whose content is the sheet's own placeholder for "nothing here".
_PLACEHOLDERS = {"", "...", "'...'", "none", "nan"}

_ELEMENT_SEPARATORS = re.compile(r"[;,]")

# The element terms the vocabulary already knows, read from
# `SEED_VOCABULARY` rather than re-listed here so de-pluralisation cannot
# drift away from the vocabulary it is trying to land on.
_KNOWN_ELEMENTS = frozenset(SEED_VOCABULARY["element"])


class ImportReport:
    """What one importer run did, and — as importantly — what it would not do.

    An import over a hand-kept spreadsheet is mostly refusals: a row with no
    dataset, a row whose `DATASET` cell holds a sentence, a channel that was
    never registered. A count alone cannot be acted on, so every refusal is
    recorded as `{"ref": <the row, as a person names it>, "reason": <why>}` and
    the caller (Prompt 05's Library page, or a script) shows the list.

    `counts` is a plain dict of running totals rather than a fixed set of
    attributes because the two importers count different things and a shared
    shape with half its fields always zero says less than the keys each one
    actually used.

    Lives here rather than in `importers/__init__.py` only because this wave's
    file ownership put that file in another agent's hands; it belongs one level
    up once the importers are all landed.
    """

    def __init__(self, source, *, dry_run=False):
        self.source = source
        self.dry_run = bool(dry_run)
        self.counts = {}
        self.skips = []
        self.notes = []

    def count(self, key, n=1):
        """Add `n` to a running total, creating the key on first use."""
        self.counts[key] = self.counts.get(key, 0) + n
        return self.counts[key]

    def skip(self, ref, reason):
        """Record a row that was not imported, and why, in the source's own
        naming — `"ID 14"`, not a zero-based sheet index."""
        self.skips.append({"ref": str(ref), "reason": str(reason)})
        self.count("skipped")

    def note(self, text):
        """A remark about the run as a whole, not about one row."""
        self.notes.append(str(text))

    def as_dict(self):
        """A JSON-safe view, for a route or a log line."""
        return {"source": self.source, "dry_run": self.dry_run,
                "counts": dict(self.counts), "skips": list(self.skips),
                "notes": list(self.notes)}

    def __repr__(self):  # pragma: no cover - diagnostics only
        return "<ImportReport {} counts={} skips={}{}>".format(
            self.source, self.counts, len(self.skips),
            " DRY-RUN" if self.dry_run else "")

    # `counts` is read far more often than it is written, and a missing key is
    # an honest zero rather than a KeyError: "the importer never got as far as
    # counting sequences" and "it counted none" are the same claim to a caller.
    def __getitem__(self, key):
        return self.counts.get(key, 0)


def sequence_key_for_span(recording_id, channel, start_idx):
    """The `sequences.sequence_key` a human span gets, from its coordinates.

    Both human importers — this one and `annotations.py` — can reach the same
    span: a catalogue row becomes an `annotations` row, and the annotations
    importer then sees that row. `sequences` is `UNIQUE (sequence_key, origin)`,
    so the two must agree on the key or the same claim lands twice under two
    names. Deriving it from the coordinates rather than from either source's
    own id is what makes them agree without one importing the other.
    """
    return "human_r{}_ch{}_{}s".format(
        int(recording_id), "x" if channel is None else int(channel),
        int(start_idx))


def global_channel(pack, channel):
    """The `recordings.channel` a `(Pack, Channel)` pair names.

    `Channel` in the spreadsheet is the electrode's index *within its pack*, so
    it is 0..3 on every row and repeats across packs. Verified against the
    drop-motif seed for all 15 catalogue ids it contains.
    """
    return int(pack) * CHANNELS_PER_PACK + int(channel)


def species_for(experiment, pack):
    """The species of an `(Experiment, Pack)` pair, or None if the legend does
    not cover it.

    None is a result, not a failure: the two experiments reuse `Pack` numbers
    for different mushrooms, so a pair outside the legend cannot be inferred
    from either half alone and is left untagged.
    """
    if experiment is None or pack is None:
        return None
    try:
        pack = int(pack)
    except (TypeError, ValueError):
        return None
    return SPECIES_BY_EXPERIMENT_PACK.get((str(experiment).strip(), pack))


def _clean(value):
    """A cell's text, or None if it is empty or one of the sheet's placeholders."""
    if value is None:
        return None
    text = str(value).strip()
    return None if text.lower() in _PLACEHOLDERS else text


def split_elements(text):
    """`Elements` -> a list of vocabulary terms, in the sheet's own order.

    Split on both `;` and `,` because the sheet uses both, sometimes in one
    cell. Case and internal spaces are normalised (`"tonic bursting"` ->
    `tonic_bursting`), the known typos are corrected, and a trailing plural
    `s` is dropped **only** when the singular is a term the vocabulary already
    has — so `sharkfins` becomes `sharkfin` while `furrycaterpillars`, which
    the vocabulary does not know in either form, is left exactly as written
    rather than mangled.
    """
    text = _clean(text)
    if not text:
        return []
    out = []
    for piece in _ELEMENT_SEPARATORS.split(text):
        term = _clean(piece)
        if not term:
            continue
        term = re.sub(r"\s+", "_", term.strip().lower())
        term = ELEMENT_TYPOS.get(term, term)
        if term not in _KNOWN_ELEMENTS and term.endswith("s") \
                and term[:-1] in _KNOWN_ELEMENTS:
            term = term[:-1]
        if term not in out:
            out.append(term)
    return out


def read_catalogue_rows(xlsx_path):
    """The spreadsheet as a list of dicts, header-keyed, in sheet order.

    `openpyxl` in `read_only` + `data_only` mode, so formulas arrive as their
    cached values and the workbook is never rewritten. Fully blank lines are
    dropped; everything else is handed on exactly as typed, because this
    function's job is to read the sheet and `import_catalogue`'s job is to
    decide what a cell means.

    `Working.registration.kinds._check_catalogue` parses the same sheet, but it
    parses it *into a registration `Report`* — it takes a `Candidate` and a
    `Report` and writes checks onto them — so there is no plain parse to call.
    The two agree on what a row is; the type coercion lives there because a
    registration check must say which cell was unparseable, and stays out of
    here because an importer wants the raw value to refuse on.
    """
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        raw = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if not raw:
        return []
    header = [str(h).strip() if h is not None else "" for h in raw[0]]
    rows = []
    for line in raw[1:]:
        if all(v is None for v in line):
            continue
        rows.append({col: value for col, value in zip(header, line) if col})
    return rows


def _candidate_source_files(dataset):
    """The `recordings.source_file` names a `DATASET` cell could mean.

    The cell holds a short dataset name (`M2_aug`, `M2`) and the recordings
    table holds the `.mat` file those channels were materialised from
    (`M2_aug_concat_fs1.mat`). One cell holds a sentence instead — the
    caller tests for that before asking.
    """
    name = dataset.strip()
    if name.lower().endswith(".mat"):
        return [name]
    return ["{}_concat_fs1.mat".format(name), "{}.mat".format(name)]


def _bind_row(conn, row):
    """Resolve one spreadsheet row to `(recording_row, start_idx, end_idx)`.

    Returns `(binding, None)` on success and `(None, reason)` on a refusal, so
    every way a hand-kept row can fail to be a span comes back as a sentence a
    person can act on instead of an exception three frames up.
    """
    dataset = _clean(row.get("DATASET"))
    if not dataset:
        return None, "no DATASET: the row does not say which recording it is in"
    if " " in dataset:
        return None, ("the DATASET cell holds a note, not a dataset name: "
                      "{!r}".format(dataset))

    candidates = _candidate_source_files(dataset)
    if any(os.path.basename(c) == HELD_OUT_RECORDING_FILE for c in candidates):
        return None, ("{} is held out for evaluation and is refused on every "
                      "route".format(HELD_OUT_RECORDING_FILE))

    pack, channel = row.get("Pack"), row.get("Channel")
    if pack is None or channel is None:
        return None, "no Pack/Channel: the row does not say which electrode it is on"
    try:
        global_ch = global_channel(pack, channel)
    except (TypeError, ValueError):
        return None, "Pack/Channel are not numbers: {!r}/{!r}".format(pack, channel)

    recording = None
    for source_file in candidates:
        recording = q.get_recording(conn, source_file, global_ch)
        if recording is not None:
            break
    if recording is None:
        return None, ("no registered recording for dataset {!r} channel {} "
                      "(Pack {} x {} + Channel {})".format(
                          dataset, global_ch, pack, CHANNELS_PER_PACK, channel))

    start_h, stop_h = row.get("StartTime_h"), row.get("StopTime_h")
    if start_h is None:
        return None, "no StartTime_h: the row has no span"
    if stop_h is None:
        return None, "no StopTime_h: the row has no span"
    try:
        fs = float(recording["fs"] or 1.0)
        start_idx = round(float(start_h) * 3600.0 * fs)
        end_idx = round(float(stop_h) * 3600.0 * fs)
    except (TypeError, ValueError):
        return None, "StartTime_h/StopTime_h are not numbers: {!r}/{!r}".format(
            start_h, stop_h)
    if end_idx <= start_idx:
        return None, ("StopTime_h {} is not after StartTime_h {} — that is not "
                      "a span".format(stop_h, start_h))

    return (recording, start_idx, end_idx), None


def _note_for(row):
    """The annotation's note: `sequence_structure`, with `Notes` appended.

    Matches what the 31 already-imported rows carry, `" | "`-joined, so a
    re-import of a row that is already there produces the same text rather than
    a near-miss that looks like a change.
    """
    parts = [_clean(row.get("sequence_structure")), _clean(row.get("Notes"))]
    parts = [p for p in parts if p]
    return " | ".join(parts) if parts else None


def import_catalogue(conn, xlsx_path, *, dry_run=False, progress=None):
    """Import `signal_catalog.xlsx` into `annotations` and `sequences`.

    Every row that binds becomes one human annotation (`source =
    'excel_catalog'`) plus one human sequence with `needs_extraction = 1`, and
    every row that does not becomes an entry in `report.skips` naming the row
    as the sheet names it and saying why. Six of the live sheet's 37 rows were
    never imported; one of them (ID 19) binds and is backfilled by this, and
    the other five do not and are reported.

    Idempotent on `(recording_id, start_idx, end_idx, source)` — see the module
    docstring for why that key and not another.

    `dry_run=True` counts exactly what a real run would write and writes
    nothing, so the Library page can show the plan before it is applied.
    `progress(done, total, what)` is called once per row, matching the
    callback shape the rest of the core uses.

    Returns an `ImportReport`.
    """
    report = ImportReport(str(xlsx_path), dry_run=dry_run)
    # Seeded so a caller reading `counts["annotations"]` on a run that imported
    # nothing gets an honest zero rather than a KeyError.
    for key in ("annotations", "sequences", "already_imported", "skipped",
                "n_events_parsed", "elements", "species_unknown"):
        report.counts.setdefault(key, 0)
    rows = read_catalogue_rows(xlsx_path)
    total = len(rows)

    for i, row in enumerate(rows, start=1):
        ref = "ID {}".format(row.get("ID_Number"))
        binding, reason = _bind_row(conn, row)
        if binding is None:
            report.skip(ref, reason)
            if progress is not None:
                progress(i, total, ref)
            continue

        recording, start_idx, end_idx = binding
        recording_id = recording["id"]
        note = _note_for(row)

        if q.annotation_exists(conn, recording_id, start_idx, end_idx,
                               CATALOGUE_SOURCE):
            report.count("already_imported")
            if progress is not None:
                progress(i, total, ref)
            continue

        # Deferred so the count is right in a dry run without the parser being
        # a second definition of "how many events does the note claim".
        from Working.library.importers.annotations import parse_event_count
        n_events = parse_event_count(note)

        report.count("annotations")
        report.count("sequences")
        if n_events is not None:
            report.count("n_events_parsed")

        elements = split_elements(row.get("Elements"))
        species = species_for(row.get("Experiment"), row.get("Pack"))
        report.count("elements", len(elements))
        if species is None:
            report.count("species_unknown")

        if not dry_run:
            annotation_id = q.insert_annotation(
                conn, recording_id, start_idx, end_idx, "interesting",
                source=CATALOGUE_SOURCE, note=note,
                status=_clean(row.get("STATUS")), event_count=n_events,
                commit=False,
            )
            _tag(conn, annotation_id, "provenance", [CATALOGUE_SOURCE])
            _tag(conn, annotation_id, "element", elements)
            if species is not None:
                _tag(conn, annotation_id, "species", [species])
            _insert_sequence(
                conn, recording_id=recording_id, channel=recording["channel"],
                start_idx=start_idx, end_idx=end_idx, n_events=n_events,
                source_kind="catalogue", source_store=str(xlsx_path),
                source_ref=str(row.get("ID_Number")), annotation_id=annotation_id,
            )

        if progress is not None:
            progress(i, total, ref)

    if not dry_run:
        conn.commit()
    return report


def _tag(conn, annotation_id, category, values):
    """Attach vocabulary terms to an annotation, creating unknown ones.

    `get_or_create_term` rather than `set_annotation_tags`, deliberately: the
    catalogue's `Elements` column is semi-controlled and holds terms the seed
    vocabulary has never seen (`bad_news`, `cx_spike_structure`). An importer
    that raised on those would refuse the row over a label, and one that
    dropped them would lose the observation; creating the term keeps both the
    row and the word the researcher used.
    """
    for value in values:
        tag_id = get_or_create_term(conn, category, value)
        conn.execute(
            "INSERT OR IGNORE INTO annotation_tags (annotation_id, tag_id) "
            "VALUES (?, ?)", (annotation_id, tag_id),
        )


def _insert_sequence(conn, *, recording_id, channel, start_idx, end_idx,
                     n_events, source_kind, source_store, source_ref,
                     annotation_id):
    """Write one human sequence, or leave the existing one alone.

    `INSERT OR IGNORE` on `UNIQUE (sequence_key, origin)` is what makes the two
    human importers able to see the same span without writing it twice —
    whichever runs first owns the row.

    `content_hash` stays NULL: a sequence's hash is over its gap profile, and a
    sequence with `needs_extraction = 1` has no resolved events and therefore
    no gaps to hash (§3.3). A hash of its whole span would be the hash of a
    single waveform, which is precisely the conflation the sequence tables
    exist to avoid.
    """
    conn.execute(
        """INSERT OR IGNORE INTO sequences
               (sequence_key, origin, recording_id, channel, start_idx,
                end_idx, n_events, needs_extraction, source_kind,
                source_store, source_ref, annotation_id, content_hash,
                created_at)
           VALUES (?, 'human', ?, ?, ?, ?, ?, 1, ?, ?, ?, ?, NULL,
                   datetime('now'))""",
        (sequence_key_for_span(recording_id, channel, start_idx),
         recording_id, channel, start_idx, end_idx, n_events,
         source_kind, source_store, source_ref, annotation_id),
    )
